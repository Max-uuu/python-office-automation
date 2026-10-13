"""
Extract tables from PDF files (e.g. invoices) into Excel.

What it does
------------
1. Opens every .pdf in the input folder
2. Finds the tables on each page and extracts them
3. Writes one Excel file:
   - All Items : every table row, with the source file and page number
   - one sheet per PDF, so each document can be checked on its own

Works with PDFs that have selectable text. Scanned PDFs (images) need OCR first.

Usage
-----
    python pdf_tables_to_excel.py --input sample_input --output tables.xlsx
"""

import argparse
import re
import sys
from pathlib import Path

import pandas as pd
import pdfplumber
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def to_number(value):
    """Turn '1,200.00' or '$45' into a number; leave other text as it is."""
    if value is None:
        return value
    text = str(value).strip()
    cleaned = re.sub(r"[,$€£¥\s]", "", text)
    if re.fullmatch(r"-?\d+(\.\d+)?", cleaned):
        return float(cleaned)
    return text


def extract_tables(pdf_path: Path) -> pd.DataFrame:
    rows = []
    header = None
    with pdfplumber.open(pdf_path) as pdf:
        for page_no, page in enumerate(pdf.pages, start=1):
            for table in page.extract_tables():
                if not table:
                    continue
                first = [str(c or "").strip() for c in table[0]]
                # Treat the first row as the header; repeat headers on later pages are skipped.
                if header is None:
                    header = first
                body = table[1:] if first == header else table
                for r in body:
                    if not any(c and str(c).strip() for c in r):
                        continue
                    rows.append([page_no] + [to_number(c) for c in r])

    if header is None:
        return pd.DataFrame()
    df = pd.DataFrame(rows, columns=["Page"] + header)
    df.insert(0, "Source File", pdf_path.name)
    return df


def style_sheet(ws):
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
    for idx, col in enumerate(ws.columns, start=1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col) + 3
        ws.column_dimensions[get_column_letter(idx)].width = min(max(width, 10), 45)
    ws.freeze_panes = "A2"


def safe_sheet_name(name: str) -> str:
    name = re.sub(r"[\[\]\:\*\?\/\\]", "_", Path(name).stem)
    return name[:31] or "Sheet"


def main():
    parser = argparse.ArgumentParser(description="Extract tables from PDFs into Excel.")
    parser.add_argument("--input", default="sample_input", help="Folder with .pdf files")
    parser.add_argument("--output", default="tables.xlsx", help="Output Excel file")
    args = parser.parse_args()

    pdfs = sorted(Path(args.input).glob("*.pdf"))
    if not pdfs:
        sys.exit(f"No .pdf files found in {args.input}")

    per_file = {}
    for pdf in pdfs:
        df = extract_tables(pdf)
        if df.empty:
            print(f"  no table found: {pdf.name}")
            continue
        per_file[pdf.name] = df
        print(f"  {pdf.name}: {len(df)} row(s)")

    if not per_file:
        sys.exit("No tables extracted.")

    all_items = pd.concat(per_file.values(), ignore_index=True)
    with pd.ExcelWriter(args.output, engine="openpyxl") as writer:
        all_items.to_excel(writer, sheet_name="All Items", index=False)
        style_sheet(writer.sheets["All Items"])
        for name, df in per_file.items():
            sheet = safe_sheet_name(name)
            df.to_excel(writer, sheet_name=sheet, index=False)
            style_sheet(writer.sheets[sheet])

    print(f"\nPDFs processed : {len(per_file)}")
    print(f"Total rows     : {len(all_items)}")
    print(f"Saved to       : {args.output}")


if __name__ == "__main__":
    main()
